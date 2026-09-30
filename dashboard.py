"""Violation review dashboard for enforcement staff (Excel workbook).

Sheets
    Overview          live KPI tiles driven by formulas over the Decisions sheet
    Review Queue      manual review cases in priority order with input columns
    Priority Alerts   stolen vehicle and Amber Alert hits for immediate action
    Auto Issue Batch  evidence complete citations ready for issuance
    Corridor Hotspots enforcement value and speeding rate per corridor and band
    Model Benchmark   speed and accuracy scorecard per perception model
    Data Integrity    integrity issues found in the source export
    Decisions         every capture with its decision (privacy safe fields only)
    Charts            key figures

The workbook never contains a raw plate string: only the masked plate and the
keyed plate token are exported.
"""

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import config as cfg

HEADER_FILL = PatternFill("solid", start_color="1F4E79")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", bold=True, size=16, color="1F4E79")
SUBTITLE_FONT = Font(name="Arial", italic=True, size=10, color="595959")
KPI_LABEL_FONT = Font(name="Arial", size=10, color="595959")
KPI_VALUE_FONT = Font(name="Arial", bold=True, size=18, color="1F4E79")
INPUT_FILL = PatternFill("solid", start_color="FFF2CC")
TILE_FILL = PatternFill("solid", start_color="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DATE_FORMAT = "dd mmm yyyy"

DECISION_COLUMNS = [
    ("capture_id", "Capture ID", 16),
    ("observation_date", "Observation Date", 14),
    ("time_of_day", "Time Band", 15),
    ("corridor_name", "Corridor", 22),
    ("district_zone", "Zone", 22),
    ("plate_masked", "Plate (Masked)", 15),
    ("plate_token", "Plate Token", 20),
    ("plate_state", "State", 7),
    ("vehicle_type", "Vehicle", 11),
    ("violation_type", "Recorded Violation", 20),
    ("rules_violation_detail", "Rules Violation", 20),
    ("decision", "Decision", 15),
    ("rules_fine_usd", "Fine (USD)", 11),
    ("speed_detected_mph", "Speed (mph)", 11),
    ("legal_speed_limit_mph", "Limit (mph)", 11),
    ("speed_excess_mph", "Excess (mph)", 11),
    ("ocr_confidence_score", "OCR Conf", 10),
    ("detection_confidence", "Detection Conf", 13),
    ("image_quality_score", "Image Quality", 12),
    ("model_name", "Model", 12),
    ("decision_reasons", "Reasons", 55),
    ("review_priority_score", "Priority Score", 12),
    ("retention_class", "Retention Class", 16),
    ("purge_date", "Purge Date", 13),
]

NUMBER_FORMATS = {
    "rules_fine_usd": "$#,##0",
    "speed_detected_mph": "0.0",
    "legal_speed_limit_mph": "0",
    "speed_excess_mph": "0.0",
    "ocr_confidence_score": "0.000",
    "detection_confidence": "0.000",
    "image_quality_score": "0.0",
    "review_priority_score": "0.00",
    "observation_date": DATE_FORMAT,
    "purge_date": DATE_FORMAT,
}


def _clean_value(value):
    if value is None:
        return None
    try:
        import pandas as pd
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(value, "to_pydatetime"):
        return value.to_pydatetime()
    if hasattr(value, "item"):
        return value.item()
    return value


def _write_table(ws, frame, columns, start_row=1, number_formats=None):
    """Write a styled table and return the last row written."""
    number_formats = number_formats or {}
    for c, (_, header, width) in enumerate(columns, start=1):
        cell = ws.cell(row=start_row, column=c, value=header)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[start_row].height = 30
    keys = [key for key, _, _ in columns]
    for r, row in enumerate(frame[keys].itertuples(index=False), start=start_row + 1):
        for c, (key, value) in enumerate(zip(keys, row), start=1):
            cell = ws.cell(row=r, column=c, value=_clean_value(value))
            cell.font = BODY_FONT
            if key in number_formats:
                cell.number_format = number_formats[key]
    last_row = start_row + len(frame)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if len(frame):
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(columns))}{last_row}"
    return last_row


def _simple_columns(frame, width=16):
    return [(c, c.replace("_", " ").title().replace("Pct", "(%)").replace("Usd", "(USD)").replace("Ms", "(ms)"), width)
            for c in frame.columns]


def _overview(ws, n_decisions):
    ws.sheet_view.showGridLines = False
    ws["B2"] = "Violation Review Dashboard"
    ws["B2"].font = TITLE_FONT
    ws["B3"] = "Automated License Plate Recognition and Traffic Violation Detection"
    ws["B3"].font = SUBTITLE_FONT

    rng = f"Decisions!$L$2:$L${n_decisions + 1}"
    fine = f"Decisions!$M$2:$M${n_decisions + 1}"
    tiles = [
        ("Captures processed", f"=COUNTA(Decisions!$A$2:$A${n_decisions + 1})", "#,##0", 5, 2),
        ("Priority alerts", f'=COUNTIF({rng},"{cfg.DECISION_PRIORITY_ALERT}")', "#,##0", 5, 4),
        ("Auto issue citations", f'=COUNTIF({rng},"{cfg.DECISION_AUTO_ISSUE}")', "#,##0", 5, 6),
        ("Manual review cases", f'=COUNTIF({rng},"{cfg.DECISION_MANUAL_REVIEW}")', "#,##0", 5, 8),
        ("Rejected (false positive)", f'=COUNTIF({rng},"{cfg.DECISION_REJECT}")', "#,##0", 8, 2),
        ("Auto issue value (USD)", f'=SUMIF({rng},"{cfg.DECISION_AUTO_ISSUE}",{fine})', "$#,##0", 8, 4),
        ("Value awaiting review (USD)", f'=SUMIF({rng},"{cfg.DECISION_MANUAL_REVIEW}",{fine})', "$#,##0", 8, 6),
        ("Review hours at current policy", f"=H6*{cfg.REVIEW_MINUTES_PER_CASE}/60", "#,##0.0", 8, 8),
    ]
    for label, formula, fmt, row, col in tiles:
        label_cell = ws.cell(row=row, column=col, value=label)
        label_cell.font = KPI_LABEL_FONT
        label_cell.fill = TILE_FILL
        value_cell = ws.cell(row=row + 1, column=col, value=formula)
        value_cell.font = KPI_VALUE_FONT
        value_cell.fill = TILE_FILL
        value_cell.number_format = fmt
        value_cell.alignment = Alignment(horizontal="left")
    for col in [2, 4, 6, 8]:
        ws.column_dimensions[get_column_letter(col)].width = 30
    for col in [1, 3, 5, 7]:
        ws.column_dimensions[get_column_letter(col)].width = 3

    notes = [
        "How to use this workbook",
        "1. Work the Review Queue from the top: cases are ordered by priority score (fine value and evidence strength).",
        "2. Record an outcome in the yellow Reviewer Decision column (Confirm, Dismiss or Escalate) and add notes if needed.",
        "3. Priority Alerts are hotlist hits (stolen vehicle, Amber Alert) and must be actioned before the queue.",
        "4. Auto Issue Batch lists citations that passed every evidence gate and can be issued without review.",
        "5. Plates are shown masked with a keyed token; the raw plate is held only in the evidence store.",
        "Assumptions",
        f"Speed tolerance {cfg.SPEED_TOLERANCE_MPH:g} mph above the posted limit at the capture point.",
        f"Evidence gates: OCR confidence at least {cfg.EVIDENCE_GATES['ocr_confidence_min']:.2f}, detection confidence at least "
        f"{cfg.EVIDENCE_GATES['detection_confidence_min']:.2f}, image quality at least {cfg.EVIDENCE_GATES['image_quality_min']:g}, plate matched to registry.",
        f"Review effort {cfg.REVIEW_MINUTES_PER_CASE:g} minutes per case. Fine schedule set in src/alpr_analytics/config.py.",
        "Red light cases require signal phase data for automatic issue; until that feed is connected they route to review.",
    ]
    for i, text in enumerate(notes, start=12):
        cell = ws.cell(row=i, column=2, value=text)
        cell.font = Font(name="Arial", bold=text in ("How to use this workbook", "Assumptions"), size=10)
    legend = ws.cell(row=12 + len(notes) + 1, column=2, value="Yellow cells in the Review Queue are reviewer inputs")
    legend.fill = INPUT_FILL
    legend.font = BODY_FONT


def build_dashboard(decisions, corridor, hotspot, benchmark, integrity, figure_paths, path=None):
    """Create the review workbook and return its path."""
    path = path or cfg.DASHBOARD_DIR / "violation_review_dashboard.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    _overview(wb.active, len(decisions))
    wb.active.title = "Overview"

    queue = decisions[decisions["decision"].eq(cfg.DECISION_MANUAL_REVIEW)].sort_values(
        "review_priority_score", ascending=False).copy()
    queue.insert(0, "queue_rank", range(1, len(queue) + 1))
    queue["reviewer_decision"] = None
    queue["reviewer_notes"] = None
    queue_cols = [("queue_rank", "Rank", 7)] + DECISION_COLUMNS[:22] + [
        ("reviewer_decision", "Reviewer Decision", 16), ("reviewer_notes", "Reviewer Notes", 30)]
    ws = wb.create_sheet("Review Queue")
    last = _write_table(ws, queue, queue_cols, number_formats=NUMBER_FORMATS)
    input_columns = [i for i, (key, _, _) in enumerate(queue_cols, start=1)
                     if key in ("reviewer_decision", "reviewer_notes")]
    for col_index in input_columns:
        for r in range(2, last + 1):
            ws.cell(row=r, column=col_index).fill = INPUT_FILL
    decision_letter = get_column_letter(min(input_columns))
    validation = DataValidation(type="list", formula1='"Confirm,Dismiss,Escalate"', allow_blank=True)
    ws.add_data_validation(validation)
    validation.add(f"{decision_letter}2:{decision_letter}{max(last, 2)}")

    alerts = decisions[decisions["decision"].eq(cfg.DECISION_PRIORITY_ALERT)].sort_values("observation_date")
    _write_table(wb.create_sheet("Priority Alerts"), alerts, DECISION_COLUMNS, number_formats=NUMBER_FORMATS)

    auto = decisions[decisions["decision"].eq(cfg.DECISION_AUTO_ISSUE)].sort_values(
        "rules_fine_usd", ascending=False)
    _write_table(wb.create_sheet("Auto Issue Batch"), auto, DECISION_COLUMNS, number_formats=NUMBER_FORMATS)

    ws = wb.create_sheet("Corridor Hotspots")
    last = _write_table(ws, corridor, _simple_columns(corridor, 18),
                        number_formats={"fine_value_usd": "$#,##0"})
    heat = hotspot.reset_index()
    heat_start = last + 3
    ws.cell(row=heat_start, column=1, value="Speeding rate (%) by corridor and time band").font = Font(
        name="Arial", bold=True, size=11, color="1F4E79")
    heat_last = _write_table(ws, heat, _simple_columns(heat, 16), start_row=heat_start + 1)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = None
    ref = f"B{heat_start + 2}:{get_column_letter(heat.shape[1])}{heat_last}"
    ws.conditional_formatting.add(ref, ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                      end_type="max", end_color="2E75B6"))

    _write_table(wb.create_sheet("Model Benchmark"), benchmark, _simple_columns(benchmark, 14))
    _write_table(wb.create_sheet("Data Integrity"), integrity, _simple_columns(integrity, 22)[:1] + [
        ("records", "Records", 12), ("share_pct", "Share (%)", 12)])

    _write_table(wb.create_sheet("Decisions"), decisions, DECISION_COLUMNS, number_formats=NUMBER_FORMATS)

    ws = wb.create_sheet("Charts")
    ws.sheet_view.showGridLines = False
    row = 1
    for figure in figure_paths:
        image = XLImage(str(figure))
        scale = 720 / image.width
        image.width, image.height = int(image.width * scale), int(image.height * scale)
        ws.add_image(image, f"B{row + 1}")
        row += int(image.height / 20) + 3

    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path
